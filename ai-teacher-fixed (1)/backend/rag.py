"""
Minimal, dependency-light RAG pipeline.

Design choice for the hackathon MVP: instead of a hosted vector DB + neural
embeddings (extra API cost, extra service to run, extra thing to explain to
judges), retrieval uses scikit-learn's TF-IDF + cosine similarity, computed
in-memory per session. It's free, needs no network call, runs in
milliseconds for a few hundred chunks, and is easy to explain on stage:
"we score which parts of your notes best match the concept we're teaching."
Swap-in path to a real vector DB (e.g. pgvector / Chroma / Pinecone) is noted
in the README for teams who want to extend this after the hackathon.
"""
from typing import List, Dict, Any, Optional
import io
import re

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


def extract_text(filename: str, raw_bytes: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return _extract_pdf(raw_bytes)
    if name.endswith(".docx"):
        return _extract_docx(raw_bytes)
    if name.endswith(".pptx"):
        return _extract_pptx(raw_bytes)
    # fall back: treat as plain text
    return raw_bytes.decode("utf-8", errors="ignore")


def extract_pages(filename: str, raw_bytes: bytes) -> List[Dict[str, Any]]:
    """Like extract_text, but preserves page numbers where the format has
    them (PDF). Returns [{"text": ..., "page": int|None}, ...] — one entry
    per page for PDFs, one entry with page=None for everything else. This is
    what makes "Source: Page 37"-style citations possible instead of just
    "Source: notes.pdf".
    """
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw_bytes))
        return [{"text": (page.extract_text() or ""), "page": i + 1} for i, page in enumerate(reader.pages)]
    return [{"text": extract_text(filename, raw_bytes), "page": None}]


def _extract_pdf(raw_bytes: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(raw_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(raw_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_pptx(raw_bytes: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(raw_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
    return "\n".join(lines)


def chunk_text(text: str, target_words: int = 120, overlap_words: int = 20) -> List[str]:
    """Sliding-window chunking on whitespace-tokenized words.

    Simple on purpose: word-count windows are predictable, cheap, and good
    enough for TF-IDF retrieval, which cares about term overlap, not about
    respecting sentence boundaries perfectly.
    """
    text = re.sub(r"\s+", " ", text).strip()
    words = text.split(" ")
    if not words or words == [""]:
        return []
    chunks = []
    step = max(target_words - overlap_words, 1)
    for start in range(0, len(words), step):
        piece = words[start:start + target_words]
        if len(piece) < 15 and chunks:
            # too small a tail chunk -> merge into previous
            chunks[-1] = chunks[-1] + " " + " ".join(piece)
            break
        chunks.append(" ".join(piece))
        if start + target_words >= len(words):
            break
    return chunks


def chunk_pages(pages: List[Dict[str, Any]], target_words: int = 120, overlap_words: int = 20) -> List[Dict[str, Any]]:
    """Chunk each page independently so every chunk keeps its page number.
    Falls back gracefully for non-paginated sources (page=None everywhere).
    """
    out = []
    for p in pages:
        for chunk in chunk_text(p["text"], target_words=target_words, overlap_words=overlap_words):
            out.append({"text": chunk, "page": p["page"]})
    return out


class Retriever:
    """Fits a TF-IDF index over a session's chunks and retrieves top-k matches."""

    def __init__(self, chunks: List[str], metadata: Optional[List[Dict[str, Any]]] = None):
        self.chunks = chunks
        # metadata[i] should describe chunks[i], e.g. {"source": "notes.pdf", "page": 3}
        self.metadata = metadata or [{} for _ in chunks]
        self._vectorizer = None
        self._matrix = None
        if chunks:
            self._vectorizer = TfidfVectorizer(stop_words="english", max_features=4096)
            self._matrix = self._vectorizer.fit_transform(chunks)

    def is_ready(self) -> bool:
        return self._matrix is not None and self._matrix.shape[0] > 0

    def top_k(self, query: str, k: int = 4) -> List[str]:
        if not self.is_ready():
            return []
        q_vec = self._vectorizer.transform([query])
        sims = cosine_similarity(q_vec, self._matrix)[0]
        ranked = sims.argsort()[::-1][:k]
        return [self.chunks[i] for i in ranked if sims[i] > 0.0]

    def top_k_with_meta(self, query: str, k: int = 4) -> List[Dict[str, Any]]:
        """Same ranking as top_k, but returns {"text", "source", "page"} so the
        teach turn can cite where the explanation came from.
        """
        if not self.is_ready():
            return []
        q_vec = self._vectorizer.transform([query])
        sims = cosine_similarity(q_vec, self._matrix)[0]
        ranked = sims.argsort()[::-1][:k]
        return [
            {"text": self.chunks[i], **self.metadata[i]}
            for i in ranked if sims[i] > 0.0
        ]
